from flask import Flask, render_template, request, redirect, url_for, send_file
import pandas as pd
import shutil
from docx import Document
import os
import zipfile
from fpdf import FPDF
from datetime import datetime

app = Flask(__name__)

# Function to generate shortage letters
def generate_letters(df):
    low_attendance_students = df[df["ATTENDANCE"] < 75]
    path = "template.docx"
    output_folder = "static/shortage_letters"

    if not os.path.exists(output_folder):
        os.makedirs(output_folder)

    for index, row in low_attendance_students.iterrows():
        name = row["STUDENT_NAME"]
        roll = row["ROLL_NUMBER"]
        year = row["YEAR_SEM"]
        dep = row["DEPARTMENT"]
        percentage = row["ATTENDANCE"]
        address = row["ADDRESS"]
        tuition_fee = row["Tuition_fees"]
        hostel_fee = row["Hostel_fees"]
        transport_fee = row["Transport_fees"]
        report_date = datetime.today().strftime('%Y-%m-%d')

        copy_path = os.path.join(output_folder, f"{roll}.docx")
        shutil.copyfile(path, copy_path)
        copy_doc = Document(copy_path)

        data_to_replace = {
            "__name": name,
            "__roll": roll,
            "__year": year,
            "__dep": dep,
            "__percent": str(percentage),
            "__address": address,
            "__tu": str(tuition_fee),
            "__ho": str(hostel_fee),
            "__tr": str(transport_fee),
            "__date": report_date
        }

        for paragraph in copy_doc.paragraphs:
            for key, value in data_to_replace.items():
                if key in paragraph.text:
                    paragraph.text = paragraph.text.replace(key, value)

        copy_doc.save(copy_path)

# Function to generate PDF
def generate_pdf(students, list_type, year_sem, section):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    pdf.cell(200, 10, txt=f"{list_type.capitalize()} List ({year_sem} - {section})", ln=True, align='C')
    pdf.ln(10)
    
    pdf.set_font("Arial", size=10)
    pdf.cell(40, 10, "Roll Number", border=1)
    pdf.cell(100, 10, "Student Name", border=1)
    pdf.cell(30, 10, "Attendance", border=1)
    pdf.ln()

    for student in students:
        pdf.cell(40, 10, student['ROLL_NUMBER'], border=1)
        pdf.cell(100, 10, student['STUDENT_NAME'], border=1)
        pdf.cell(30, 10, str(round(float(student['ATTENDANCE']), 2)), border=1)
        pdf.ln()
    
    file_name = f"{list_type}_list_for_{year_sem}_and_{section}"
    file_name = file_name.replace(" ", "_").replace("/", "-")
    pdf_output = f"static/{file_name}.pdf"
    pdf.output(pdf_output)
    return pdf_output

# Login page
@app.route('/', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        if request.form['username'] == 'admin' and request.form['password'] == 'password':
            return redirect(url_for('dashboard'))
        else:
            return render_template('login.html', error=True)
    return render_template('login.html', error=False)


# Dashboard page
@app.route('/dashboard')
def dashboard():
    return render_template('dashboard.html')

# Upload attendance file page
@app.route('/upload', methods=['GET', 'POST'])
def upload_file():
    if request.method == 'POST':
        file = request.files['file']
        if file.filename != '':
            file.save("attendance.xlsx")
            df = pd.read_excel("attendance.xlsx")
            generate_letters(df)
            return redirect(url_for('dashboard'))
    return render_template('upload.html')

# View attendance page
@app.route('/attendance', methods=['GET', 'POST'])
def view_attendance():
    if request.method == 'POST':
        year_sem = request.form['year_sem']
        section = request.form['section']
        df = pd.read_excel("attendance.xlsx")
        students = df[(df['YEAR_SEM'] == year_sem) & (df['SECTION'] == section)].to_dict(orient='records')
        return render_template('attendance.html', students=students, year_sem=year_sem, section=section)
    df = pd.read_excel("attendance.xlsx")
    year_sems = df['YEAR_SEM'].unique()
    sections = df['SECTION'].unique()
    return render_template('view_attendance.html', year_sems=year_sems, sections=sections)

# Overall reports
@app.route('/overall_reports', methods=['POST'])
def overall_reports():
    year_sem = request.form.get('year_sem')
    section = request.form.get('section')
    df = pd.read_excel("attendance.xlsx")
    students = df[(df['YEAR_SEM'] == year_sem) & (df['SECTION'] == section) & (df['ATTENDANCE'] < 75)].to_dict(orient='records')
    return render_template('overall_reports.html', students=students, year_sem=year_sem, section=section)

# Condonation list
@app.route('/condonation_list', methods=['POST'])
def condonation_list():
    year_sem = request.form.get('year_sem')
    section = request.form.get('section')
    df = pd.read_excel("attendance.xlsx")
    students = df[(df['YEAR_SEM'] == year_sem) & (df['SECTION'] == section) & (df['ATTENDANCE'] >= 65) & (df['ATTENDANCE'] < 75)].to_dict(orient='records')
    return render_template('condonation_list.html', students=students, year_sem=year_sem, section=section)

# Detained list
@app.route('/detained_list', methods=['POST'])
def detained_list():
    year_sem = request.form.get('year_sem')
    section = request.form.get('section')
    df = pd.read_excel("attendance.xlsx")
    students = df[(df['YEAR_SEM'] == year_sem) & (df['SECTION'] == section) & (df['ATTENDANCE'] < 65)].to_dict(orient='records')
    return render_template('detained_list.html', students=students, year_sem=year_sem, section=section)

# Generate reports
@app.route('/generate_reports', methods=['POST'])
def generate_reports():
    year_sem = request.form.get('year_sem')
    section = request.form.get('section')
    df = pd.read_excel("attendance.xlsx")
    students = df[(df['YEAR_SEM'] == year_sem) & (df['SECTION'] == section) & (df['ATTENDANCE'] < 75)].to_dict(orient='records')
    return render_template('generate_reports.html', students=students, year_sem=year_sem, section=section)

# Download list
@app.route('/download_list/<list_type>', methods=['POST'])
def download_list(list_type):
    year_sem = request.form.get('year_sem')
    section = request.form.get('section')
    df = pd.read_excel("attendance.xlsx")
    if list_type == 'overall':
        students = df[(df['YEAR_SEM'] == year_sem) & (df['SECTION'] == section) & (df['ATTENDANCE'] < 75)].to_dict(orient='records')
    elif list_type == 'condonation':
        students = df[(df['YEAR_SEM'] == year_sem) & (df['SECTION'] == section) & (df['ATTENDANCE'] >= 65) & (df['ATTENDANCE'] < 75)].to_dict(orient='records')
    elif list_type == 'detained':
        students = df[(df['YEAR_SEM'] == year_sem) & (df['SECTION'] == section) & (df['ATTENDANCE'] < 65)].to_dict(orient='records')
    else:
        students = []

    pdf_path = generate_pdf(students, list_type, year_sem, section)
    return send_file(pdf_path, as_attachment=True)

# Download all generated reports
@app.route('/download_reports', methods=['POST'])
def download_reports():
    year_sem = request.form.get('year_sem')
    section = request.form.get('section')
    output_folder = "static/shortage_letters"
    file_name = f"Reports for {year_sem} and {section}"
    file_name = file_name.replace(" ", "_").replace("/", "-")
    zip_path = f"static/{file_name}.zip"

    # Create a zip file containing all letters
    with zipfile.ZipFile(zip_path, 'w') as zipf:
        for root, dirs, files in os.walk(output_folder):
            for file in files:
                zipf.write(os.path.join(root, file), os.path.relpath(os.path.join(root, file), output_folder))

    # Serve the zip file for download
    return send_file(zip_path, as_attachment=True)

if __name__ == '__main__':
    app.run(debug=True)
